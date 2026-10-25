"""Requisitos / cláusulas derivados de uma Norma (submódulo Requisitos).

CRUD por norma + import/export (CSV sempre; DOCX/XLSX quando a lib estiver presente).
Vínculos de rastreabilidade e sugestão por IA são acrescentados nas Fases 4 e 5
(editam este mesmo arquivo). Aditivo: não altera módulos existentes.
"""
import csv
import io
import json
import re
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import ConfigDict, Field

from app.auth import get_current_user, require_papel
from app.db import get_connection
from app.routers.llm_acessos import scope
from app.security import SafeModel

router = APIRouter(prefix="/api/requisitos", tags=["requisitos"])

editor_dep = require_papel("admin", "editor")

CATEGORIAS = ("normativo", "interno", "cliente", "legal")
OBRIGATORIEDADE = ("obrigatorio", "recomendado", "informativo")
ESTADOS = ("ativo", "revogado")
PRIORIDADES = ("", "must", "should", "could", "wont")  # MoSCoW (IREB cap. 4)
TIPO_VINCULO = ("satisfaz", "verifica", "implementa", "mitiga", "viola", "relacionado", "outro")  # requisito↔controle (IREB cap. 6); viola/relacionado p/ Issue/RNC
TIPO_RELACAO = ("deriva_de", "refina", "depende", "conflita", "duplica", "variante_de")  # requisito↔requisito
ATRIBUTO_TIPOS = ("texto", "numero", "data", "lista", "booleano")

DDL = """CREATE TABLE IF NOT EXISTS requisitos (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 norma_id INTEGER NOT NULL REFERENCES normas(id) ON DELETE CASCADE,
 codigo_clausula TEXT DEFAULT '',
 titulo TEXT NOT NULL,
 texto TEXT DEFAULT '',
 categoria TEXT NOT NULL DEFAULT 'normativo',
 obrigatoriedade TEXT NOT NULL DEFAULT 'obrigatorio',
 estado TEXT NOT NULL DEFAULT 'ativo',
 ordem INTEGER DEFAULT 0,
 criado_por_id INTEGER REFERENCES usuarios(id),
 criado_em TEXT NOT NULL,
 atualizado_em TEXT
);
CREATE INDEX IF NOT EXISTS idx_requisitos_norma ON requisitos(norma_id, ordem);
CREATE TRIGGER IF NOT EXISTS requisitos_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisitos WHERE empresa_id=OLD.id; END;
CREATE TRIGGER IF NOT EXISTS requisitos_usuario_delete BEFORE DELETE ON usuarios
BEGIN UPDATE requisitos SET criado_por_id=NULL WHERE criado_por_id=OLD.id; END;
CREATE TABLE IF NOT EXISTS requisito_atributo_defs (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 chave TEXT NOT NULL,
 rotulo TEXT NOT NULL,
 tipo TEXT NOT NULL DEFAULT 'texto',
 opcoes TEXT DEFAULT '',
 obrigatorio INTEGER NOT NULL DEFAULT 0,
 ordem INTEGER DEFAULT 0,
 ativo INTEGER NOT NULL DEFAULT 1,
 criado_em TEXT NOT NULL,
 UNIQUE(empresa_id, chave)
);
CREATE TABLE IF NOT EXISTS requisito_atributo_valores (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 requisito_id INTEGER NOT NULL REFERENCES requisitos(id) ON DELETE CASCADE,
 def_id INTEGER NOT NULL REFERENCES requisito_atributo_defs(id) ON DELETE CASCADE,
 valor TEXT DEFAULT '',
 UNIQUE(requisito_id, def_id)
);
CREATE TABLE IF NOT EXISTS requisito_relacoes (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 origem_id INTEGER NOT NULL REFERENCES requisitos(id) ON DELETE CASCADE,
 destino_id INTEGER NOT NULL REFERENCES requisitos(id) ON DELETE CASCADE,
 tipo TEXT NOT NULL DEFAULT 'deriva_de',
 observacao TEXT DEFAULT '',
 criado_por_id INTEGER REFERENCES usuarios(id),
 criado_em TEXT NOT NULL,
 UNIQUE(origem_id, destino_id, tipo)
);
CREATE TRIGGER IF NOT EXISTS req_atrib_defs_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisito_atributo_defs WHERE empresa_id=OLD.id; END;
CREATE TRIGGER IF NOT EXISTS req_relacoes_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisito_relacoes WHERE empresa_id=OLD.id; END;
CREATE TRIGGER IF NOT EXISTS req_relacoes_usuario_delete BEFORE DELETE ON usuarios
BEGIN UPDATE requisito_relacoes SET criado_por_id=NULL WHERE criado_por_id=OLD.id; END;
CREATE TABLE IF NOT EXISTS requisito_mudancas (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 requisito_id INTEGER NOT NULL REFERENCES requisitos(id) ON DELETE CASCADE,
 descricao TEXT NOT NULL,
 situacao TEXT NOT NULL DEFAULT 'proposta',
 solicitante_id INTEGER REFERENCES usuarios(id),
 aprovador_id INTEGER REFERENCES usuarios(id),
 decisao_obs TEXT DEFAULT '',
 criado_em TEXT NOT NULL,
 decidido_em TEXT
);
CREATE TRIGGER IF NOT EXISTS req_mudancas_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisito_mudancas WHERE empresa_id=OLD.id; END;
CREATE TRIGGER IF NOT EXISTS req_mudancas_usuario_delete BEFORE DELETE ON usuarios
BEGIN UPDATE requisito_mudancas SET solicitante_id=NULL WHERE solicitante_id=OLD.id;
 UPDATE requisito_mudancas SET aprovador_id=NULL WHERE aprovador_id=OLD.id; END;
CREATE TABLE IF NOT EXISTS requisito_versoes (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 requisito_id INTEGER NOT NULL REFERENCES requisitos(id) ON DELETE CASCADE,
 numero INTEGER NOT NULL,
 snapshot_json TEXT NOT NULL,
 origem TEXT DEFAULT 'manual',
 criado_por_id INTEGER REFERENCES usuarios(id),
 criado_em TEXT NOT NULL,
 UNIQUE(requisito_id, numero)
);
CREATE TRIGGER IF NOT EXISTS req_versoes_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisito_versoes WHERE empresa_id=OLD.id; END;
CREATE TRIGGER IF NOT EXISTS req_versoes_usuario_delete BEFORE DELETE ON usuarios
BEGIN UPDATE requisito_versoes SET criado_por_id=NULL WHERE criado_por_id=OLD.id; END;
"""

# Classes (seções) que agrupam requisitos dentro de cada norma (Norma → Classe → Requisito).
CLASSES_DDL = """CREATE TABLE IF NOT EXISTS requisito_classes (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
 norma_id INTEGER NOT NULL REFERENCES normas(id) ON DELETE CASCADE,
 codigo TEXT DEFAULT '',
 nome TEXT NOT NULL,
 ordem INTEGER DEFAULT 0,
 criado_em TEXT NOT NULL,
 UNIQUE(norma_id, nome)
);
CREATE INDEX IF NOT EXISTS idx_req_classes_norma ON requisito_classes(norma_id, ordem);
CREATE TRIGGER IF NOT EXISTS req_classes_empresa_delete BEFORE DELETE ON empresas
BEGIN DELETE FROM requisito_classes WHERE empresa_id=OLD.id; END;
"""


class RequisitoIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    norma_id: int
    codigo_clausula: str = Field(default="", max_length=60)
    titulo: str = Field(min_length=1, max_length=300)
    texto: str = Field(default="", max_length=8000)
    categoria: str = "normativo"
    obrigatoriedade: str = "obrigatorio"
    estado: str = "ativo"
    prioridade: str = ""
    variante: str = Field(default="", max_length=120)
    fonte: str = Field(default="", max_length=500)
    justificativa: str = Field(default="", max_length=4000)
    classe_id: int | None = None
    ordem: int = 0
    custo_atual: float | None = Field(default=None, ge=0)
    custo_projetado: float | None = Field(default=None, ge=0)
    volume: float | None = Field(default=None, ge=0)


class RequisitoUpdate(SafeModel):
    model_config = ConfigDict(extra="forbid")
    codigo_clausula: str | None = Field(default=None, max_length=60)
    titulo: str | None = Field(default=None, max_length=300)
    texto: str | None = Field(default=None, max_length=8000)
    categoria: str | None = None
    obrigatoriedade: str | None = None
    estado: str | None = None
    prioridade: str | None = None
    variante: str | None = Field(default=None, max_length=120)
    fonte: str | None = Field(default=None, max_length=500)
    justificativa: str | None = Field(default=None, max_length=4000)
    classe_id: int | None = None
    ordem: int | None = None
    custo_atual: float | None = Field(default=None, ge=0)
    custo_projetado: float | None = Field(default=None, ge=0)
    volume: float | None = Field(default=None, ge=0)


def now():
    return datetime.now(timezone.utc).isoformat()


def _valida_enums(categoria, obrigatoriedade, estado):
    if categoria not in CATEGORIAS:
        raise HTTPException(status_code=422, detail="Categoria inválida")
    if obrigatoriedade not in OBRIGATORIEDADE:
        raise HTTPException(status_code=422, detail="Obrigatoriedade inválida")
    if estado not in ESTADOS:
        raise HTTPException(status_code=422, detail="Estado inválido")


def _norma_da_empresa(conn, norma_id, company):
    row = conn.execute("SELECT id FROM normas WHERE id=? AND empresa_id=?", (norma_id, company)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Norma não encontrada")


def _valida_classe(conn, classe_id, norma_id, company):
    """Classe (opcional) precisa pertencer à mesma norma e empresa do requisito."""
    if classe_id is None:
        return
    row = conn.execute("SELECT id FROM requisito_classes WHERE id=? AND norma_id=? AND empresa_id=?",
                       (classe_id, norma_id, company)).fetchone()
    if row is None:
        raise HTTPException(status_code=422, detail="Classe inválida para esta norma")


def req_to_dict(r):
    return dict(r)


def fetch_req(conn, req_id, company):
    row = conn.execute("SELECT * FROM requisitos WHERE id=? AND empresa_id=?", (req_id, company)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Requisito não encontrado")
    return row


@router.get("")
def list_requisitos(norma_id: int | None = None, empresa_id: int | None = None, categoria: str = "",
                    estado: str = "", prioridade: str = "", q: str = "", user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    # Sem norma_id: visão global dos requisitos da empresa (todas as normas).
    sql = ("SELECT r.*, n.codigo AS norma_codigo, n.titulo AS norma_titulo, c.nome AS classe_nome "
           "FROM requisitos r JOIN normas n ON n.id=r.norma_id "
           "LEFT JOIN requisito_classes c ON c.id=r.classe_id WHERE r.empresa_id=?")
    params = [company]
    if norma_id is not None:
        sql += " AND r.norma_id=?"
        params.append(norma_id)
    if categoria:
        sql += " AND r.categoria=?"
        params.append(categoria)
    if estado:
        sql += " AND r.estado=?"
        params.append(estado)
    if prioridade:
        sql += " AND r.prioridade=?"
        params.append(prioridade)
    if q:
        sql += " AND (r.codigo_clausula LIKE ? OR r.titulo LIKE ? OR r.texto LIKE ?)"
        like = f"%{q}%"
        params.extend([like, like, like])
    sql += " ORDER BY n.codigo, r.ordem, r.id"
    with get_connection() as conn:
        if norma_id is not None:
            _norma_da_empresa(conn, norma_id, company)
        return [req_to_dict(r) for r in conn.execute(sql, params).fetchall()]


@router.post("", status_code=201)
def create_requisito(body: RequisitoIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    _valida_enums(body.categoria, body.obrigatoriedade, body.estado)
    if body.prioridade not in PRIORIDADES:
        raise HTTPException(status_code=422, detail="Prioridade inválida")
    instante = now()
    with get_connection() as conn:
        _norma_da_empresa(conn, body.norma_id, company)
        _valida_classe(conn, body.classe_id, body.norma_id, company)
        # Todo requisito nasce de uma norma: fonte e justificativa vazias herdam a norma de origem.
        fonte = body.fonte
        justif = body.justificativa
        if not (fonte or "").strip() or not (justif or "").strip():
            nrow = conn.execute("SELECT identificador_externo, titulo FROM normas WHERE id=?", (body.norma_id,)).fetchone()
            src = ((nrow["identificador_externo"] or nrow["titulo"]) if nrow else "") or ""
            src = src.strip()
            if src and not (fonte or "").strip():
                fonte = src
            if src and not (justif or "").strip():
                justif = f"Requisito decorrente da norma {src}"
        cur = conn.execute(
            "INSERT INTO requisitos (empresa_id,norma_id,codigo_clausula,titulo,texto,categoria,obrigatoriedade,estado,prioridade,variante,fonte,justificativa,classe_id,ordem,custo_atual,custo_projetado,volume,criado_por_id,criado_em,atualizado_em) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (company, body.norma_id, body.codigo_clausula, body.titulo, body.texto, body.categoria,
             body.obrigatoriedade, body.estado, body.prioridade, body.variante, fonte, justif, body.classe_id,
             body.ordem, body.custo_atual, body.custo_projetado, body.volume, user["id"], instante, instante),
        )
        conn.commit()
        return req_to_dict(conn.execute("SELECT * FROM requisitos WHERE id=?", (cur.lastrowid,)).fetchone())


@router.put("/{req_id}")
def update_requisito(req_id: int, body: RequisitoUpdate, empresa_id: int | None = None,
                     user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    instante = now()
    with get_connection() as conn:
        row = fetch_req(conn, req_id, company)
        categoria = body.categoria if body.categoria is not None else row["categoria"]
        obrigatoriedade = body.obrigatoriedade if body.obrigatoriedade is not None else row["obrigatoriedade"]
        estado = body.estado if body.estado is not None else row["estado"]
        _valida_enums(categoria, obrigatoriedade, estado)
        prioridade = body.prioridade if body.prioridade is not None else (row["prioridade"] or "")
        if prioridade not in PRIORIDADES:
            raise HTTPException(status_code=422, detail="Prioridade inválida")
        campos = dict(
            codigo_clausula=body.codigo_clausula if body.codigo_clausula is not None else row["codigo_clausula"],
            titulo=body.titulo if body.titulo is not None else row["titulo"],
            texto=body.texto if body.texto is not None else row["texto"],
            categoria=categoria, obrigatoriedade=obrigatoriedade, estado=estado, prioridade=prioridade,
            variante=body.variante if body.variante is not None else (row["variante"] or ""),
            fonte=body.fonte if body.fonte is not None else (row["fonte"] if "fonte" in row.keys() else ""),
            justificativa=body.justificativa if body.justificativa is not None else (row["justificativa"] if "justificativa" in row.keys() else ""),
            ordem=body.ordem if body.ordem is not None else row["ordem"],
            custo_atual=body.custo_atual if body.custo_atual is not None else row["custo_atual"],
            custo_projetado=body.custo_projetado if body.custo_projetado is not None else row["custo_projetado"],
            volume=body.volume if body.volume is not None else row["volume"],
            classe_id=(body.classe_id if "classe_id" in body.model_fields_set
                       else (row["classe_id"] if "classe_id" in row.keys() else None)),
        )
        _valida_classe(conn, campos["classe_id"], row["norma_id"], company)
        conn.execute(
            "UPDATE requisitos SET codigo_clausula=?,titulo=?,texto=?,categoria=?,obrigatoriedade=?,estado=?,prioridade=?,variante=?,fonte=?,justificativa=?,ordem=?,custo_atual=?,custo_projetado=?,volume=?,classe_id=?,atualizado_em=? WHERE id=?",
            (*campos.values(), instante, req_id),
        )
        conn.commit()
        return req_to_dict(conn.execute("SELECT * FROM requisitos WHERE id=?", (req_id,)).fetchone())


@router.delete("/{req_id}", status_code=204)
def delete_requisito(req_id: int, empresa_id: int | None = None, user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        conn.execute("DELETE FROM requisitos WHERE id=?", (req_id,))
        conn.commit()


# --- IREB cap. 3: atributos (schema por empresa) ---

admin_dep = require_papel("admin")


class AtributoDefIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    chave: str = Field(min_length=1, max_length=60)
    rotulo: str = Field(min_length=1, max_length=120)
    tipo: str = "texto"
    opcoes: list[str] = Field(default_factory=list, max_length=100)
    obrigatorio: bool = False
    ordem: int = 0
    ativo: bool = True


def _def_to_dict(d):
    out = dict(d)
    out["obrigatorio"] = bool(out["obrigatorio"])
    out["ativo"] = bool(out["ativo"])
    try:
        out["opcoes"] = json.loads(out.get("opcoes") or "[]")
    except (ValueError, TypeError):
        out["opcoes"] = []
    return out


@router.get("/atributos/defs")
def list_atributo_defs(empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM requisito_atributo_defs WHERE empresa_id=? ORDER BY ordem,id", (company,)
        ).fetchall()
        return [_def_to_dict(r) for r in rows]


@router.post("/atributos/defs", status_code=201)
def create_atributo_def(body: AtributoDefIn, user: dict = Depends(admin_dep)):
    import sqlite3
    company = scope(user, body.empresa_id)
    if body.tipo not in ATRIBUTO_TIPOS:
        raise HTTPException(status_code=422, detail="Tipo de atributo inválido")
    if body.tipo == "lista" and not body.opcoes:
        raise HTTPException(status_code=422, detail="Atributo do tipo lista exige opções")
    with get_connection() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO requisito_atributo_defs (empresa_id,chave,rotulo,tipo,opcoes,obrigatorio,ordem,ativo,criado_em) "
                "VALUES (?,?,?,?,?,?,?,?,?)",
                (company, body.chave, body.rotulo, body.tipo, json.dumps(body.opcoes, ensure_ascii=False),
                 int(body.obrigatorio), body.ordem, int(body.ativo), now()),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="Já existe um atributo com essa chave")
        conn.commit()
        return _def_to_dict(conn.execute("SELECT * FROM requisito_atributo_defs WHERE id=?", (cur.lastrowid,)).fetchone())


@router.put("/atributos/defs/{def_id}")
def update_atributo_def(def_id: int, body: AtributoDefIn, user: dict = Depends(admin_dep)):
    company = scope(user, body.empresa_id)
    if body.tipo not in ATRIBUTO_TIPOS:
        raise HTTPException(status_code=422, detail="Tipo de atributo inválido")
    if body.tipo == "lista" and not body.opcoes:
        raise HTTPException(status_code=422, detail="Atributo do tipo lista exige opções")
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM requisito_atributo_defs WHERE id=? AND empresa_id=?", (def_id, company)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Atributo não encontrado")
        conn.execute(
            "UPDATE requisito_atributo_defs SET rotulo=?,tipo=?,opcoes=?,obrigatorio=?,ordem=?,ativo=? WHERE id=?",
            (body.rotulo, body.tipo, json.dumps(body.opcoes, ensure_ascii=False), int(body.obrigatorio),
             body.ordem, int(body.ativo), def_id),
        )
        conn.commit()
        return _def_to_dict(conn.execute("SELECT * FROM requisito_atributo_defs WHERE id=?", (def_id,)).fetchone())


@router.delete("/atributos/defs/{def_id}", status_code=204)
def delete_atributo_def(def_id: int, empresa_id: int | None = None, user: dict = Depends(admin_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        row = conn.execute("SELECT id FROM requisito_atributo_defs WHERE id=? AND empresa_id=?", (def_id, company)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Atributo não encontrado")
        conn.execute("DELETE FROM requisito_atributo_defs WHERE id=?", (def_id,))
        conn.commit()


def _atributos_do_requisito(conn, company, req_id):
    defs = conn.execute(
        "SELECT * FROM requisito_atributo_defs WHERE empresa_id=? AND ativo=1 ORDER BY ordem,id", (company,)
    ).fetchall()
    vals = {v["def_id"]: v["valor"] for v in conn.execute(
        "SELECT def_id,valor FROM requisito_atributo_valores WHERE requisito_id=?", (req_id,)).fetchall()}
    out = []
    for d in defs:
        dd = _def_to_dict(d)
        dd["valor"] = vals.get(d["id"], "")
        out.append(dd)
    return out


def _valida_valor_atributo(tipo, opcoes, valor):
    if valor == "":
        return True
    if tipo == "numero":
        try:
            float(valor)
        except (ValueError, TypeError):
            return False
    elif tipo == "booleano":
        return valor in ("0", "1", "true", "false")
    elif tipo == "lista":
        return valor in (opcoes or [])
    return True


class ValoresIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    valores: dict[str, str] = Field(default_factory=dict)


@router.put("/{req_id}/atributos")
def set_atributos(req_id: int, body: ValoresIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        defs = {str(d["id"]): _def_to_dict(d) for d in conn.execute(
            "SELECT * FROM requisito_atributo_defs WHERE empresa_id=? AND ativo=1", (company,)).fetchall()}
        for def_id, valor in body.valores.items():
            d = defs.get(str(def_id))
            if d is None:
                raise HTTPException(status_code=422, detail=f"Atributo {def_id} inválido")
            if not _valida_valor_atributo(d["tipo"], d["opcoes"], valor):
                raise HTTPException(status_code=422, detail=f"Valor inválido para '{d['rotulo']}'")
        for d in defs.values():
            if d["obrigatorio"] and not (body.valores.get(str(d["id"])) or "").strip():
                raise HTTPException(status_code=422, detail=f"Atributo obrigatório '{d['rotulo']}' não informado")
        for def_id, valor in body.valores.items():
            conn.execute(
                "INSERT INTO requisito_atributo_valores (requisito_id,def_id,valor) VALUES (?,?,?) "
                "ON CONFLICT(requisito_id,def_id) DO UPDATE SET valor=excluded.valor",
                (req_id, int(def_id), valor),
            )
        conn.commit()
        return {"atributos": _atributos_do_requisito(conn, company, req_id)}


# --- IREB cap. 6: relações requisito↔requisito ---

class RelacaoIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    destino_id: int
    tipo: str = "deriva_de"
    observacao: str = Field(default="", max_length=2000)


class ObsIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    observacao: str = Field(default="", max_length=2000)


@router.get("/{req_id}/relacoes")
def list_relacoes(req_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        rows = conn.execute(
            "SELECT rr.*, r.codigo_clausula AS destino_clausula, r.titulo AS destino_titulo "
            "FROM requisito_relacoes rr JOIN requisitos r ON r.id=rr.destino_id WHERE rr.origem_id=? ORDER BY rr.id",
            (req_id,),
        ).fetchall()
        return [dict(r) for r in rows]


@router.post("/{req_id}/relacoes", status_code=201)
def add_relacao(req_id: int, body: RelacaoIn, user: dict = Depends(editor_dep)):
    import sqlite3
    company = scope(user, body.empresa_id)
    if body.tipo not in TIPO_RELACAO:
        raise HTTPException(status_code=422, detail="Tipo de relação inválido")
    if body.destino_id == req_id:
        raise HTTPException(status_code=422, detail="Um requisito não se relaciona consigo mesmo")
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        if conn.execute("SELECT id FROM requisitos WHERE id=? AND empresa_id=?", (body.destino_id, company)).fetchone() is None:
            raise HTTPException(status_code=404, detail="Requisito de destino não encontrado")
        try:
            cur = conn.execute(
                "INSERT INTO requisito_relacoes (empresa_id,origem_id,destino_id,tipo,observacao,criado_por_id,criado_em) "
                "VALUES (?,?,?,?,?,?,?)",
                (company, req_id, body.destino_id, body.tipo, body.observacao, user["id"], now()),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="Relação já existe")
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_relacoes WHERE id=?", (cur.lastrowid,)).fetchone())


@router.put("/{req_id}/relacoes/{relacao_id}")
def editar_relacao(req_id: int, relacao_id: int, body: ObsIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        row = conn.execute(
            "SELECT id FROM requisito_relacoes WHERE id=? AND origem_id=? AND empresa_id=?", (relacao_id, req_id, company)
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Relação não encontrada")
        conn.execute("UPDATE requisito_relacoes SET observacao=? WHERE id=?", (body.observacao, relacao_id))
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_relacoes WHERE id=?", (relacao_id,)).fetchone())


@router.delete("/{req_id}/relacoes/{relacao_id}", status_code=204)
def del_relacao(req_id: int, relacao_id: int, empresa_id: int | None = None, user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        row = conn.execute(
            "SELECT id FROM requisito_relacoes WHERE id=? AND origem_id=? AND empresa_id=?", (relacao_id, req_id, company)
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Relação não encontrada")
        conn.execute("DELETE FROM requisito_relacoes WHERE id=?", (relacao_id,))
        conn.commit()


def _linhas_do_arquivo(nome, data):
    """Retorna lista de dicts {codigo_clausula,titulo,texto,categoria,obrigatoriedade}."""
    ext = (nome or "").lower().rsplit(".", 1)[-1] if "." in (nome or "") else ""
    linhas = []
    if ext == "csv":
        texto = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(texto))
        for r in reader:
            low = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
            titulo = low.get("titulo") or low.get("título") or low.get("requisito") or ""
            if not titulo:
                continue
            linhas.append(dict(
                codigo_clausula=low.get("codigo_clausula") or low.get("clausula") or low.get("cláusula") or "",
                titulo=titulo, texto=low.get("texto") or "",
                categoria=low.get("categoria") or "normativo",
                obrigatoriedade=low.get("obrigatoriedade") or "obrigatorio",
            ))
    elif ext == "docx":
        try:
            from docx import Document
        except ImportError:
            raise HTTPException(status_code=400, detail="Suporte a DOCX indisponível (python-docx ausente)")
        doc = Document(io.BytesIO(data))
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                m = re.match(r"^\s*([0-9]+(?:\.[0-9]+)*)[\)\.\s-]+(.*)$", t)
                linhas.append(dict(codigo_clausula=m.group(1) if m else "", titulo=(m.group(2) if m else t)[:300],
                                   texto="", categoria="normativo", obrigatoriedade="obrigatorio"))
    elif ext == "xlsx":
        try:
            import openpyxl
        except ImportError:
            raise HTTPException(status_code=400, detail="Suporte a XLSX indisponível (openpyxl ausente)")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            header = [str(c or "").strip().lower() for c in rows[0]]
            idx = {name: header.index(name) for name in header}
            def get(row, *names):
                for n in names:
                    if n in idx and idx[n] < len(row):
                        return str(row[idx[n]] or "").strip()
                return ""
            for row in rows[1:]:
                titulo = get(row, "titulo", "título", "requisito")
                if not titulo:
                    continue
                linhas.append(dict(codigo_clausula=get(row, "codigo_clausula", "clausula", "cláusula"),
                                   titulo=titulo, texto=get(row, "texto"),
                                   categoria=get(row, "categoria") or "normativo",
                                   obrigatoriedade=get(row, "obrigatoriedade") or "obrigatorio"))
    else:
        raise HTTPException(status_code=400, detail="Formato não suportado (use CSV, DOCX ou XLSX)")
    return linhas


@router.post("/importar", status_code=201)
def importar(norma_id: int = Form(...), empresa_id: int | None = Form(default=None),
             file: UploadFile = File(...), user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    data = file.file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Arquivo vazio")
    from app.security import sanitize_text
    linhas = _linhas_do_arquivo(file.filename or "", data)
    instante = now()
    criados = 0
    with get_connection() as conn:
        _norma_da_empresa(conn, norma_id, company)
        base = conn.execute("SELECT COALESCE(MAX(ordem),0) FROM requisitos WHERE norma_id=?", (norma_id,)).fetchone()[0]
        for i, ln in enumerate(linhas, start=1):
            categoria = ln["categoria"] if ln["categoria"] in CATEGORIAS else "normativo"
            obrig = ln["obrigatoriedade"] if ln["obrigatoriedade"] in OBRIGATORIEDADE else "obrigatorio"
            conn.execute(
                "INSERT INTO requisitos (empresa_id,norma_id,codigo_clausula,titulo,texto,categoria,obrigatoriedade,estado,ordem,criado_por_id,criado_em,atualizado_em) "
                "VALUES (?,?,?,?,?,?,?,'ativo',?,?,?,?)",
                (company, norma_id, sanitize_text(ln["codigo_clausula"])[:60], sanitize_text(ln["titulo"])[:300],
                 sanitize_text(ln["texto"])[:8000], categoria, obrig, base + i, user["id"], instante, instante),
            )
            criados += 1
        conn.commit()
    return {"importados": criados}


@router.get("/exportar")
def exportar(norma_id: int, empresa_id: int | None = None, formato: str = "csv",
             user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        _norma_da_empresa(conn, norma_id, company)
        rows = conn.execute(
            "SELECT codigo_clausula,titulo,texto,categoria,obrigatoriedade,estado FROM requisitos WHERE norma_id=? ORDER BY ordem,id",
            (norma_id,),
        ).fetchall()
    cols = ["codigo_clausula", "titulo", "texto", "categoria", "obrigatoriedade", "estado"]
    if formato == "xlsx":
        try:
            import openpyxl
        except ImportError:
            raise HTTPException(status_code=400, detail="Suporte a XLSX indisponível (openpyxl ausente)")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(cols)
        for r in rows:
            ws.append([r[c] for c in cols])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="requisitos_{norma_id}.xlsx"'})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow([r[c] for c in cols])
    return StreamingResponse(io.BytesIO(buf.getvalue().encode("utf-8-sig")), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="requisitos_{norma_id}.csv"'})


# --- Vínculos de rastreabilidade (Fase 4): requisito ↔ processo/issue/rnc/acao ---

class VinculoIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    alvo_tipo: str
    alvo_id: int
    tipo_relacao: str = "satisfaz"
    observacao: str = Field(default="", max_length=2000)


@router.get("/{req_id}/vinculos")
def list_vinculos(req_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        rows = conn.execute(
            "SELECT * FROM requisito_vinculos WHERE requisito_id=? ORDER BY id", (req_id,)
        ).fetchall()
        return [dict(r) for r in rows]


@router.post("/{req_id}/vinculos", status_code=201)
def add_vinculo(req_id: int, body: VinculoIn, user: dict = Depends(editor_dep)):
    import sqlite3
    from app.routers.conformidade import ALVOS, alvo_empresa
    company = scope(user, body.empresa_id)
    if body.alvo_tipo not in ALVOS:
        raise HTTPException(status_code=422, detail="Tipo de alvo inválido")
    if body.tipo_relacao not in TIPO_VINCULO:
        raise HTTPException(status_code=422, detail="Tipo de relação inválido")
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        emp = alvo_empresa(conn, body.alvo_tipo, body.alvo_id)
        if emp is None:
            raise HTTPException(status_code=404, detail="Alvo do vínculo não encontrado")
        if emp != company:
            raise HTTPException(status_code=403, detail="O alvo pertence a outra empresa")
        try:
            cur = conn.execute(
                "INSERT INTO requisito_vinculos (empresa_id,requisito_id,alvo_tipo,alvo_id,tipo_relacao,observacao,criado_por_id,criado_em) "
                "VALUES (?,?,?,?,?,?,?,?)",
                (company, req_id, body.alvo_tipo, body.alvo_id, body.tipo_relacao, body.observacao, user["id"], now()),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="Vínculo já existe")
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_vinculos WHERE id=?", (cur.lastrowid,)).fetchone())


@router.put("/{req_id}/vinculos/{vinculo_id}")
def editar_vinculo(req_id: int, vinculo_id: int, body: ObsIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        row = conn.execute(
            "SELECT id FROM requisito_vinculos WHERE id=? AND requisito_id=? AND empresa_id=?",
            (vinculo_id, req_id, company)
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Vínculo não encontrado")
        conn.execute("UPDATE requisito_vinculos SET observacao=? WHERE id=?", (body.observacao, vinculo_id))
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_vinculos WHERE id=?", (vinculo_id,)).fetchone())


@router.delete("/{req_id}/vinculos/{vinculo_id}", status_code=204)
def del_vinculo(req_id: int, vinculo_id: int, empresa_id: int | None = None, user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        row = conn.execute(
            "SELECT id FROM requisito_vinculos WHERE id=? AND requisito_id=?", (vinculo_id, req_id)
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Vínculo não encontrado")
        conn.execute("DELETE FROM requisito_vinculos WHERE id=?", (vinculo_id,))
        conn.commit()


# --- Sugestão por IA (Fase 5): parse do texto da norma → candidatos (não persiste) ---

class SugerirIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    norma_id: int
    acesso_id: int


@router.post("/sugerir")
def sugerir_por_ia(body: SugerirIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        _norma_da_empresa(conn, body.norma_id, company)
        norma = conn.execute("SELECT documento_id FROM normas WHERE id=?", (body.norma_id,)).fetchone()
        texto = ""
        if norma["documento_id"]:
            rev = conn.execute(
                "SELECT texto_extraido FROM documento_revisoes WHERE documento_id=? AND texto_extraido<>'' ORDER BY numero DESC LIMIT 1",
                (norma["documento_id"],),
            ).fetchone()
            texto = rev["texto_extraido"] if rev else ""
        if not (texto or "").strip():
            raise HTTPException(status_code=400, detail="Sem texto lido do documento da norma. Rode a leitura da revisão no PDM antes de sugerir.")
        acesso = conn.execute(
            "SELECT * FROM llm_acessos WHERE id=? AND empresa_id=? AND ativo=1", (body.acesso_id, company)
        ).fetchone()
        if acesso is None:
            raise HTTPException(status_code=404, detail="Acesso LLM não encontrado")
        acesso = dict(acesso)
    from app.services import requisitos_ia
    from app.services.llm_client import LlmError
    try:
        sugestoes = requisitos_ia.sugerir(acesso, texto)
    except LlmError as e:
        raise HTTPException(status_code=502, detail=e.message)
    return {"sugestoes": sugestoes}


def _alvo_label(conn, tipo, alvo_id):
    if tipo == "processo":
        r = conn.execute("SELECT nome FROM processos WHERE id=?", (alvo_id,)).fetchone()
        return r["nome"] if r else f"Processo {alvo_id}"
    if tipo == "issue":
        r = conn.execute("SELECT codigo FROM issues WHERE id=?", (alvo_id,)).fetchone()
        return r["codigo"] if r else f"Issue {alvo_id}"
    if tipo == "rnc":
        r = conn.execute("SELECT codigo FROM rncs WHERE id=?", (alvo_id,)).fetchone()
        return r["codigo"] if r else f"RNC {alvo_id}"
    if tipo == "acao":
        return f"Ação {alvo_id}"
    if tipo == "familia":
        r = conn.execute("SELECT nome FROM produto_familias WHERE id=?", (alvo_id,)).fetchone()
        return (f"Família: {r['nome']}") if r else f"Família {alvo_id}"
    if tipo == "grupo":
        r = conn.execute("SELECT nome FROM produto_grupos WHERE id=?", (alvo_id,)).fetchone()
        return (f"Grupo: {r['nome']}") if r else f"Grupo {alvo_id}"
    if tipo == "produto":
        r = conn.execute("SELECT codigo, nome FROM produtos WHERE id=?", (alvo_id,)).fetchone()
        return (f"Produto: {(r['codigo'] + ' — ') if r['codigo'] else ''}{r['nome']}") if r else f"Produto {alvo_id}"
    return f"{tipo} {alvo_id}"


def _nome(conn, uid):
    if uid is None:
        return None
    r = conn.execute("SELECT nome FROM usuarios WHERE id=?", (uid,)).fetchone()
    return r["nome"] if r else None


# --- Change request de requisito (IREB cap. 5) ---

class MudancaIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    descricao: str = Field(min_length=1, max_length=4000)


class DecisaoIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    situacao: str
    decisao_obs: str = Field(default="", max_length=2000)


def _mud_to_dict(conn, r):
    d = dict(r)
    d["solicitante_nome"] = _nome(conn, r["solicitante_id"])
    d["aprovador_nome"] = _nome(conn, r["aprovador_id"])
    return d


@router.get("/{req_id}/mudancas")
def list_mudancas(req_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        rows = conn.execute("SELECT * FROM requisito_mudancas WHERE requisito_id=? ORDER BY id DESC", (req_id,)).fetchall()
        return [_mud_to_dict(conn, r) for r in rows]


@router.post("/{req_id}/mudancas", status_code=201)
def propor_mudanca(req_id: int, body: MudancaIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        cur = conn.execute(
            "INSERT INTO requisito_mudancas (empresa_id,requisito_id,descricao,situacao,solicitante_id,criado_em) "
            "VALUES (?,?,?,'proposta',?,?)",
            (company, req_id, body.descricao, user["id"], now()),
        )
        conn.commit()
        return _mud_to_dict(conn, conn.execute("SELECT * FROM requisito_mudancas WHERE id=?", (cur.lastrowid,)).fetchone())


@router.post("/{req_id}/mudancas/{mid}/decidir")
def decidir_mudanca(req_id: int, mid: int, body: DecisaoIn, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    if body.situacao not in ("aprovada", "rejeitada"):
        raise HTTPException(status_code=422, detail="Situação inválida (use aprovada ou rejeitada)")
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        m = conn.execute("SELECT * FROM requisito_mudancas WHERE id=? AND requisito_id=? AND empresa_id=?",
                         (mid, req_id, company)).fetchone()
        if m is None:
            raise HTTPException(status_code=404, detail="Solicitação não encontrada")
        if m["situacao"] != "proposta":
            raise HTTPException(status_code=422, detail="Solicitação já decidida")
        conn.execute("UPDATE requisito_mudancas SET situacao=?,aprovador_id=?,decisao_obs=?,decidido_em=? WHERE id=?",
                     (body.situacao, user["id"], body.decisao_obs, now(), mid))
        if body.situacao == "aprovada":
            _snapshot_requisito(conn, req_id, f"mudanca:{mid}", user["id"])
        conn.commit()
        return _mud_to_dict(conn, conn.execute("SELECT * FROM requisito_mudancas WHERE id=?", (mid,)).fetchone())


# --- Versionamento do requisito (snapshot numerado) ---
SNAP_FIELDS = ("codigo_clausula", "titulo", "texto", "categoria", "obrigatoriedade", "estado", "prioridade", "variante")


def _snapshot_requisito(conn, req_id, origem, user_id):
    r = conn.execute("SELECT * FROM requisitos WHERE id=?", (req_id,)).fetchone()
    numero = conn.execute("SELECT COALESCE(MAX(numero),0) FROM requisito_versoes WHERE requisito_id=?", (req_id,)).fetchone()[0] + 1
    snap = {k: r[k] for k in SNAP_FIELDS}
    conn.execute(
        "INSERT INTO requisito_versoes (empresa_id,requisito_id,numero,snapshot_json,origem,criado_por_id,criado_em) VALUES (?,?,?,?,?,?,?)",
        (r["empresa_id"], req_id, numero, json.dumps(snap, ensure_ascii=False), origem, user_id, now()))
    return numero


@router.get("/{req_id}/versoes")
def list_versoes(req_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        out = []
        for r in conn.execute("SELECT * FROM requisito_versoes WHERE requisito_id=? ORDER BY numero DESC", (req_id,)).fetchall():
            d = dict(r)
            d["criado_por_nome"] = _nome(conn, r["criado_por_id"])
            try:
                d["snapshot"] = json.loads(r["snapshot_json"])
            except (ValueError, TypeError):
                d["snapshot"] = {}
            d.pop("snapshot_json", None)
            out.append(d)
        return out


@router.post("/{req_id}/versoes", status_code=201)
def registrar_versao(req_id: int, empresa_id: int | None = None, user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        fetch_req(conn, req_id, company)
        numero = _snapshot_requisito(conn, req_id, "manual", user["id"])
        conn.commit()
        return {"numero": numero}


# --- Análise de impacto econômico (versão QMS do caso $150→$50 × N) ---

def _economia(ca, cp, vol):
    if ca is None or cp is None or vol is None:
        return None
    return (ca - cp) * vol


@router.get("/impacto")
def impacto(norma_id: int | None = None, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    sql = ("SELECT r.id,r.codigo_clausula,r.titulo,r.prioridade,r.custo_atual,r.custo_projetado,r.volume,"
           "n.codigo AS norma_codigo FROM requisitos r JOIN normas n ON n.id=r.norma_id WHERE r.empresa_id=?")
    params = [company]
    if norma_id is not None:
        sql += " AND r.norma_id=?"
        params.append(norma_id)
    with get_connection() as conn:
        if norma_id is not None:
            _norma_da_empresa(conn, norma_id, company)
        rows = conn.execute(sql, params).fetchall()
    itens = []
    for r in rows:
        eco = _economia(r["custo_atual"], r["custo_projetado"], r["volume"])
        d = dict(r)
        d["economia"] = eco
        itens.append(d)
    itens.sort(key=lambda x: (x["economia"] is None, -(x["economia"] or 0)))
    com_dados = [x for x in itens if x["economia"] is not None]
    resumo = dict(requisitos=len(itens), requisitos_com_dados=len(com_dados),
                  economia_total=sum(x["economia"] for x in com_dados))
    return dict(requisitos=itens, resumo=resumo)


@router.get("/arvore")
def arvore(norma_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    """Árvore de rastreabilidade de uma norma: Norma → Requisitos → vínculos e relações."""
    company = scope(user, empresa_id)
    with get_connection() as conn:
        n = conn.execute("SELECT id,codigo,titulo FROM normas WHERE id=? AND empresa_id=?", (norma_id, company)).fetchone()
        if n is None:
            raise HTTPException(status_code=404, detail="Norma não encontrada")
        classes = [dict(id=c["id"], codigo=c["codigo"], nome=c["nome"]) for c in conn.execute(
            "SELECT id,codigo,nome FROM requisito_classes WHERE norma_id=? AND empresa_id=? ORDER BY ordem,id",
            (norma_id, company)).fetchall()]
        classe_nome = {c["id"]: c["nome"] for c in classes}
        out = []
        for r in conn.execute("SELECT * FROM requisitos WHERE empresa_id=? AND norma_id=? ORDER BY ordem,id", (company, norma_id)).fetchall():
            conf = conn.execute("SELECT situacao FROM requisito_conformidade WHERE requisito_id=? AND processo_id IS NULL", (r["id"],)).fetchone()
            vincs = []
            for v in conn.execute("SELECT * FROM requisito_vinculos WHERE requisito_id=? ORDER BY id", (r["id"],)).fetchall():
                vincs.append(dict(id=v["id"], alvo_tipo=v["alvo_tipo"], alvo_id=v["alvo_id"],
                                  alvo_label=_alvo_label(conn, v["alvo_tipo"], v["alvo_id"]),
                                  observacao=(v["observacao"] if "observacao" in v.keys() else ""),
                                  tipo_relacao=(v["tipo_relacao"] if "tipo_relacao" in v.keys() else "satisfaz")))
            rels = [dict(id=rr["id"], destino_id=rr["destino_id"], destino_clausula=rr["destino_clausula"],
                         tipo=rr["tipo"], observacao=rr["observacao"])
                    for rr in conn.execute(
                        "SELECT rr.id, rr.destino_id, rr.tipo, rr.observacao, r.codigo_clausula AS destino_clausula "
                        "FROM requisito_relacoes rr JOIN requisitos r ON r.id=rr.destino_id WHERE rr.origem_id=? ORDER BY rr.id",
                        (r["id"],)).fetchall()]
            muds = [_mud_to_dict(conn, m) for m in conn.execute(
                "SELECT * FROM requisito_mudancas WHERE requisito_id=? ORDER BY id DESC", (r["id"],)).fetchall()]
            cid = r["classe_id"] if "classe_id" in r.keys() else None
            out.append(dict(id=r["id"], codigo_clausula=r["codigo_clausula"], titulo=r["titulo"],
                            categoria=r["categoria"], prioridade=(r["prioridade"] or ""),
                            obrigatoriedade=r["obrigatoriedade"], estado=r["estado"],
                            classe_id=cid, classe_nome=classe_nome.get(cid),
                            situacao=(conf["situacao"] if conf else "nao_avaliado"),
                            vinculos=vincs, relacoes=rels, mudancas=muds))
    return dict(norma=dict(id=n["id"], codigo=n["codigo"], titulo=n["titulo"]), classes=classes, requisitos=out)


@router.get("/versoes-arvore")
def versoes_arvore(norma_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    """Árvore de versionamento: por norma, cada requisito e seu histórico de versões."""
    company = scope(user, empresa_id)
    with get_connection() as conn:
        n = conn.execute("SELECT id,codigo,titulo FROM normas WHERE id=? AND empresa_id=?", (norma_id, company)).fetchone()
        if n is None:
            raise HTTPException(status_code=404, detail="Norma não encontrada")
        out = []
        for r in conn.execute("SELECT id,codigo_clausula,titulo FROM requisitos WHERE empresa_id=? AND norma_id=? ORDER BY ordem,id", (company, norma_id)).fetchall():
            versoes = [dict(numero=v["numero"], origem=v["origem"], criado_em=v["criado_em"], criado_por_nome=_nome(conn, v["criado_por_id"]))
                       for v in conn.execute("SELECT numero,origem,criado_em,criado_por_id FROM requisito_versoes WHERE requisito_id=? ORDER BY numero DESC", (r["id"],)).fetchall()]
            out.append(dict(id=r["id"], codigo_clausula=r["codigo_clausula"], titulo=r["titulo"], versoes=versoes))
    return dict(norma=dict(id=n["id"], codigo=n["codigo"], titulo=n["titulo"]), requisitos=out)


@router.get("/variantes")
def variantes(norma_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    """Consulta de variantes de uma norma: requisito base → suas variantes (relação `variante_de`)."""
    company = scope(user, empresa_id)
    with get_connection() as conn:
        n = conn.execute("SELECT id,codigo,titulo FROM normas WHERE id=? AND empresa_id=?", (norma_id, company)).fetchone()
        if n is None:
            raise HTTPException(status_code=404, detail="Norma não encontrada")
        reqs = {r["id"]: r for r in conn.execute(
            "SELECT id,codigo_clausula,titulo,variante FROM requisitos WHERE empresa_id=? AND norma_id=? ORDER BY ordem,id",
            (company, norma_id)).fetchall()}
        grupos_map = {}
        usados = set()
        for rr in conn.execute("SELECT origem_id,destino_id FROM requisito_relacoes WHERE empresa_id=? AND tipo='variante_de'", (company,)).fetchall():
            if rr["origem_id"] in reqs:  # a variante pertence a esta norma
                grupos_map.setdefault(rr["destino_id"], []).append(rr["origem_id"])
                usados.add(rr["origem_id"])
        def _r(x):
            return dict(id=x["id"], codigo_clausula=x["codigo_clausula"], titulo=x["titulo"], variante=(x["variante"] or ""))
        grupos = []
        for base_id, vids in grupos_map.items():
            base = reqs.get(base_id)
            if base is None:
                base = conn.execute("SELECT id,codigo_clausula,titulo,variante FROM requisitos WHERE id=? AND empresa_id=?", (base_id, company)).fetchone()
            if base is None:
                continue
            grupos.append(dict(base=_r(base), variantes=[_r(reqs[v]) for v in vids if v in reqs]))
        bases = set(grupos_map.keys())
        soltas = [_r(r) for r in reqs.values() if (r["variante"] or "").strip() and r["id"] not in usados and r["id"] not in bases]
    return dict(norma=dict(id=n["id"], codigo=n["codigo"], titulo=n["titulo"]), grupos=grupos, soltas=soltas)


# --- Classes de requisito por norma (Norma → Classe → Requisito) ---

class ClasseIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    norma_id: int
    codigo: str = Field(default="", max_length=60)
    nome: str = Field(min_length=1, max_length=200)
    ordem: int = 0


class ClasseUpdate(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    codigo: str | None = Field(default=None, max_length=60)
    nome: str | None = Field(default=None, max_length=200)
    ordem: int | None = None


@router.get("/classes")
def list_classes(norma_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        _norma_da_empresa(conn, norma_id, company)
        return [dict(r) for r in conn.execute(
            "SELECT * FROM requisito_classes WHERE norma_id=? AND empresa_id=? ORDER BY ordem,id",
            (norma_id, company)).fetchall()]


@router.post("/classes", status_code=201)
def create_classe(body: ClasseIn, user: dict = Depends(editor_dep)):
    import sqlite3
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        _norma_da_empresa(conn, body.norma_id, company)
        try:
            cur = conn.execute(
                "INSERT INTO requisito_classes (empresa_id,norma_id,codigo,nome,ordem,criado_em) VALUES (?,?,?,?,?,?)",
                (company, body.norma_id, body.codigo, body.nome, body.ordem, now()))
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="Classe já existe nesta norma")
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_classes WHERE id=?", (cur.lastrowid,)).fetchone())


@router.put("/classes/{classe_id}")
def update_classe(classe_id: int, body: ClasseUpdate, user: dict = Depends(editor_dep)):
    company = scope(user, body.empresa_id)
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM requisito_classes WHERE id=? AND empresa_id=?", (classe_id, company)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Classe não encontrada")
        codigo = body.codigo if body.codigo is not None else row["codigo"]
        nome = body.nome if body.nome is not None else row["nome"]
        ordem = body.ordem if body.ordem is not None else row["ordem"]
        conn.execute("UPDATE requisito_classes SET codigo=?,nome=?,ordem=? WHERE id=?", (codigo, nome, ordem, classe_id))
        conn.commit()
        return dict(conn.execute("SELECT * FROM requisito_classes WHERE id=?", (classe_id,)).fetchone())


@router.delete("/classes/{classe_id}", status_code=204)
def delete_classe(classe_id: int, empresa_id: int | None = None, user: dict = Depends(editor_dep)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        row = conn.execute("SELECT id FROM requisito_classes WHERE id=? AND empresa_id=?", (classe_id, company)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Classe não encontrada")
        conn.execute("UPDATE requisitos SET classe_id=NULL WHERE classe_id=?", (classe_id,))
        conn.execute("DELETE FROM requisito_classes WHERE id=?", (classe_id,))
        conn.commit()


# --- Clonagem / reuso (estratégia "cloning"/like-item, Grady §2.4.3) ---

class ClonarIn(SafeModel):
    model_config = ConfigDict(extra="forbid")
    empresa_id: int | None = None
    origem_norma_id: int
    destino_norma_id: int


@router.post("/clonar", status_code=201)
def clonar_requisitos(body: ClonarIn, user: dict = Depends(editor_dep)):
    """Copia todos os requisitos (campos + valores de atributos) de uma norma para outra, da mesma empresa."""
    company = scope(user, body.empresa_id)
    if body.origem_norma_id == body.destino_norma_id:
        raise HTTPException(status_code=422, detail="Origem e destino devem ser normas diferentes")
    instante = now()
    with get_connection() as conn:
        _norma_da_empresa(conn, body.origem_norma_id, company)
        _norma_da_empresa(conn, body.destino_norma_id, company)
        base = conn.execute("SELECT COALESCE(MAX(ordem),0) FROM requisitos WHERE norma_id=?", (body.destino_norma_id,)).fetchone()[0]
        origem = conn.execute("SELECT * FROM requisitos WHERE empresa_id=? AND norma_id=? ORDER BY ordem,id",
                              (company, body.origem_norma_id)).fetchall()
        criados = 0
        for i, r in enumerate(origem, start=1):
            k = r.keys()
            cur = conn.execute(
                "INSERT INTO requisitos (empresa_id,norma_id,codigo_clausula,titulo,texto,categoria,obrigatoriedade,estado,prioridade,variante,fonte,justificativa,ordem,custo_atual,custo_projetado,volume,criado_por_id,criado_em,atualizado_em) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (company, body.destino_norma_id, r["codigo_clausula"], r["titulo"], r["texto"], r["categoria"],
                 r["obrigatoriedade"], r["estado"], r["prioridade"] or "", (r["variante"] if "variante" in k else "") or "",
                 (r["fonte"] if "fonte" in k else "") or "", (r["justificativa"] if "justificativa" in k else "") or "",
                 base + i, r["custo_atual"] if "custo_atual" in k else None,
                 r["custo_projetado"] if "custo_projetado" in k else None, r["volume"] if "volume" in k else None,
                 user["id"], instante, instante),
            )
            novo_id = cur.lastrowid
            for av in conn.execute("SELECT def_id,valor FROM requisito_atributo_valores WHERE requisito_id=?", (r["id"],)).fetchall():
                conn.execute("INSERT OR IGNORE INTO requisito_atributo_valores (requisito_id,def_id,valor) VALUES (?,?,?)",
                             (novo_id, av["def_id"], av["valor"]))
            criados += 1
        conn.commit()
    return {"criados": criados, "destino_norma_id": body.destino_norma_id}


# --- Atributos de qualidade padrão (checklist de "bom requisito", IREB/Grady) ---
QUALIDADE_ATRIBUTOS = [
    ("req_verificavel", "Verificável"), ("req_nao_ambiguo", "Não ambíguo"),
    ("req_singular", "Singular/atômico"), ("req_mensuravel", "Mensurável"),
    ("req_viavel", "Viável"), ("req_rastreavel", "Rastreável"),
]


@router.post("/atributos/defs/semear-qualidade", status_code=201)
def semear_qualidade(empresa_id: int | None = None, user: dict = Depends(admin_dep)):
    """Cria (idempotente) o conjunto padrão de atributos de qualidade do requisito para a empresa."""
    company = scope(user, empresa_id)
    instante = now()
    with get_connection() as conn:
        existentes = {r["chave"] for r in conn.execute(
            "SELECT chave FROM requisito_atributo_defs WHERE empresa_id=?", (company,)).fetchall()}
        base = conn.execute("SELECT COALESCE(MAX(ordem),0) FROM requisito_atributo_defs WHERE empresa_id=?", (company,)).fetchone()[0]
        criados = 0
        for i, (chave, rotulo) in enumerate(QUALIDADE_ATRIBUTOS, start=1):
            if chave in existentes:
                continue
            conn.execute(
                "INSERT INTO requisito_atributo_defs (empresa_id,chave,rotulo,tipo,opcoes,obrigatorio,ordem,ativo,criado_em) "
                "VALUES (?,?,?,'booleano','',0,?,1,?)",
                (company, chave, rotulo, base + i, instante),
            )
            criados += 1
        conn.commit()
    return {"criados": criados}


# --- Auditoria de rastreabilidade (Grady §2.2.7): órfãos e lacunas ---

@router.get("/auditoria")
def auditoria(norma_id: int | None = None, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    sql = ("SELECT r.*, n.codigo AS norma_codigo FROM requisitos r JOIN normas n ON n.id=r.norma_id "
           "WHERE r.empresa_id=?")
    params = [company]
    with get_connection() as conn:
        if norma_id is not None:
            _norma_da_empresa(conn, norma_id, company)
            sql += " AND r.norma_id=?"
            params.append(norma_id)
        sql += " ORDER BY n.codigo, r.ordem, r.id"
        resumo = dict(total=0, sem_fonte=0, sem_justificativa=0, sem_vinculo=0, sem_verificacao=0)
        itens = []
        for r in conn.execute(sql, params).fetchall():
            k = r.keys()
            nv = conn.execute("SELECT COUNT(*) FROM requisito_vinculos WHERE requisito_id=?", (r["id"],)).fetchone()[0]
            nverif = conn.execute("SELECT COUNT(*) FROM requisito_vinculos WHERE requisito_id=? AND tipo_relacao='verifica'", (r["id"],)).fetchone()[0]
            fonte = ((r["fonte"] if "fonte" in k else "") or "").strip()
            just = ((r["justificativa"] if "justificativa" in k else "") or "").strip()
            flags = dict(sem_fonte=not fonte, sem_justificativa=not just, sem_vinculo=(nv == 0), sem_verificacao=(nverif == 0))
            resumo["total"] += 1
            for key, val in flags.items():
                if val:
                    resumo[key] += 1
            itens.append(dict(id=r["id"], codigo_clausula=r["codigo_clausula"], titulo=r["titulo"],
                              norma_codigo=r["norma_codigo"], fonte=fonte, justificativa=just, **flags))
    return dict(resumo=resumo, itens=itens)


@router.get("/rastreabilidade-export")
def rastreabilidade_export(norma_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        n = conn.execute("SELECT codigo,titulo FROM normas WHERE id=? AND empresa_id=?", (norma_id, company)).fetchone()
        if n is None:
            raise HTTPException(status_code=404, detail="Norma não encontrada")
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["Norma", "Cláusula", "Requisito", "Fonte", "Justificativa",
                    "Pais (deriva/refina)", "Vínculos", "Situação"])
        for r in conn.execute("SELECT * FROM requisitos WHERE empresa_id=? AND norma_id=? ORDER BY ordem,id",
                              (company, norma_id)).fetchall():
            k = r.keys()
            pais = conn.execute(
                "SELECT r2.codigo_clausula AS c FROM requisito_relacoes rr JOIN requisitos r2 ON r2.id=rr.destino_id "
                "WHERE rr.origem_id=? AND rr.tipo IN ('deriva_de','refina') ORDER BY rr.id", (r["id"],)).fetchall()
            vincs = conn.execute("SELECT alvo_tipo,alvo_id FROM requisito_vinculos WHERE requisito_id=? ORDER BY id", (r["id"],)).fetchall()
            conf = conn.execute("SELECT situacao FROM requisito_conformidade WHERE requisito_id=? AND processo_id IS NULL", (r["id"],)).fetchone()
            w.writerow([n["codigo"], r["codigo_clausula"] or "", r["titulo"] or "",
                        (r["fonte"] if "fonte" in k else "") or "", (r["justificativa"] if "justificativa" in k else "") or "",
                        "; ".join((p["c"] or "") for p in pais),
                        "; ".join(f'{v["alvo_tipo"]}#{v["alvo_id"]}' for v in vincs),
                        (conf["situacao"] if conf else "nao_avaliado")])
    return StreamingResponse(io.BytesIO(buf.getvalue().encode("utf-8-sig")), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="rastreabilidade_{n["codigo"]}.csv"'})


@router.get("/vinculos-por-alvo")
def vinculos_por_alvo(alvo_tipo: str, alvo_id: int, empresa_id: int | None = None,
                      user: dict = Depends(get_current_user)):
    """Consulta reversa: dado um alvo (issue/rnc/processo/acao), lista os requisitos
    vinculados e a norma de cada um. Usado pelas telas de Issue/RNC."""
    from app.routers.conformidade import ALVOS, alvo_empresa
    company = scope(user, empresa_id)
    if alvo_tipo not in ALVOS:
        raise HTTPException(status_code=422, detail="Tipo de alvo inválido")
    with get_connection() as conn:
        emp = alvo_empresa(conn, alvo_tipo, alvo_id)
        if emp is None:
            raise HTTPException(status_code=404, detail="Alvo não encontrado")
        if emp != company:
            raise HTTPException(status_code=404, detail="Alvo não encontrado")
        rows = conn.execute(
            "SELECT v.id AS vinculo_id, v.tipo_relacao, v.observacao, "
            "r.id AS requisito_id, r.codigo_clausula, r.titulo, "
            "n.id AS norma_id, n.codigo AS norma_codigo, n.titulo AS norma_titulo "
            "FROM requisito_vinculos v "
            "JOIN requisitos r ON r.id=v.requisito_id "
            "JOIN normas n ON n.id=r.norma_id "
            "WHERE v.empresa_id=? AND v.alvo_tipo=? AND v.alvo_id=? ORDER BY v.id",
            (company, alvo_tipo, alvo_id),
        ).fetchall()
        return [dict(r) for r in rows]


# Detalhe do requisito (atributos + vínculos + relações + norma). Declarado por último
# para não capturar rotas literais de um segmento como `/exportar`.
@router.get("/{req_id}")
def get_requisito(req_id: int, empresa_id: int | None = None, user: dict = Depends(get_current_user)):
    company = scope(user, empresa_id)
    with get_connection() as conn:
        d = req_to_dict(fetch_req(conn, req_id, company))
        nr = conn.execute("SELECT codigo,titulo FROM normas WHERE id=?", (d["norma_id"],)).fetchone()
        d["norma_codigo"] = nr["codigo"] if nr else ""
        d["norma_titulo"] = nr["titulo"] if nr else ""
        d["atributos"] = _atributos_do_requisito(conn, company, req_id)
        vincs = []
        for r in conn.execute("SELECT * FROM requisito_vinculos WHERE requisito_id=? ORDER BY id", (req_id,)).fetchall():
            v = dict(r)
            v["alvo_label"] = _alvo_label(conn, r["alvo_tipo"], r["alvo_id"])
            vincs.append(v)
        d["vinculos"] = vincs
        d["relacoes"] = [dict(r) for r in conn.execute(
            "SELECT rr.*, r.codigo_clausula AS destino_clausula, r.titulo AS destino_titulo "
            "FROM requisito_relacoes rr JOIN requisitos r ON r.id=rr.destino_id WHERE rr.origem_id=? ORDER BY rr.id",
            (req_id,)).fetchall()]
        d["relacoes_entrada"] = [dict(r) for r in conn.execute(
            "SELECT rr.*, r.codigo_clausula AS origem_clausula, r.titulo AS origem_titulo "
            "FROM requisito_relacoes rr JOIN requisitos r ON r.id=rr.origem_id WHERE rr.destino_id=? ORDER BY rr.id",
            (req_id,)).fetchall()]
    return d
